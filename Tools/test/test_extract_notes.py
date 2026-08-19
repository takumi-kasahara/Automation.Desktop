#!/usr/bin/env python3
"""
Test suite for extract_notes.py.

This module contains test cases for the extract_notes.py script,
which extracts notes from PowerPoint presentations.
"""

import os
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Iterator
from pathlib import Path

import pytest
from pptx import Presentation

# Add Tools directory to path so src.extract_notes can be imported
sys.path.insert(0, str(Path(__file__).parent.parent))

from src.extract_notes import build_split_paths, clean, is_sequence_template


class TestExtractNotes:
    """
    Test cases for extract_notes.py
    """

    @pytest.fixture
    def script_path(self) -> Path:
        """
        Return the path to the extract_notes.py script.
        """
        return Path(__file__).parent.parent / "src" / "extract_notes.py"

    @pytest.fixture
    def sample_pptx(self) -> Iterator[Path]:
        """
        Create a sample PowerPoint file with notes for testing.
        """
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            temp_path = Path(f.name)

        prs = Presentation()
        # Add a slide with notes
        prs.slides.add_slide(prs.slide_layouts[0])
        slide1 = prs.slides[0]
        slide1.notes_slide.notes_text_frame.text = "This is note 1 for slide 1."

        slide2 = prs.slides.add_slide(prs.slide_layouts[0])
        slide2.notes_slide.notes_text_frame.text = (
            "This is note 2 for slide 2.\nWith multiple lines."
        )

        # Slide 3 has no notes
        prs.slides.add_slide(prs.slide_layouts[0])

        slide4 = prs.slides.add_slide(prs.slide_layouts[0])
        slide4.notes_slide.notes_text_frame.text = (
            "Note with {kanji|kana} ruby and [memo] text."
        )

        prs.save(str(temp_path))
        yield temp_path
        # Cleanup
        if temp_path.exists():
            temp_path.unlink()

    @pytest.fixture
    def empty_pptx(self) -> Iterator[Path]:
        """
        Create a PowerPoint file with no notes.
        """
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            temp_path = Path(f.name)

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        # No notes added
        prs.save(str(temp_path))
        yield temp_path
        # Cleanup
        if temp_path.exists():
            temp_path.unlink()

    def run_script(
        self, script_path: Path, args: list[str], input_data: str | None = None
    ) -> subprocess.CompletedProcess[str]:
        """
        Run the extract_notes.py script with given arguments.
        """
        cmd = [sys.executable, str(script_path)] + args
        result: subprocess.CompletedProcess[str] = subprocess.run(
            cmd,
            input=input_data,
            capture_output=True,
            text=True,
            encoding="utf-8",
            check=False,
        )
        return result

    def test_extract_notes_basic(self, script_path: Path, sample_pptx: Path) -> None:
        """
        Test basic note extraction to a single output file.
        """
        # Use a non-existent output file path to avoid overwrite prompt
        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "output.txt"

            result = self.run_script(
                script_path, ["-i", str(sample_pptx), "-o", str(output_file)]
            )
            assert result.returncode == 0
            assert output_file.exists()

            content = output_file.read_text(encoding="utf-8")
            assert "This is note 1 for slide 1" in content
            assert "This is note 2 for slide 2" in content
            assert "With multiple lines" in content
            # Note 3 is empty, should not appear
            # Note 4 has ruby and memo formatting
            assert "kanji" in content  # ruby should be cleaned to kanji
            assert "memo" not in content  # memo should be removed

    def test_extract_notes_with_force(
        self, script_path: Path, sample_pptx: Path
    ) -> None:
        """
        Test force overwrite option.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "output.txt"

            # First run
            result = self.run_script(
                script_path, ["-i", str(sample_pptx), "-o", str(output_file)]
            )
            assert result.returncode == 0

            # Second run with --force should succeed without prompt
            result = self.run_script(
                script_path, ["-i", str(sample_pptx), "-o", str(output_file), "-f"]
            )
            assert result.returncode == 0

    def test_extract_notes_split_mode(
        self, script_path: Path, sample_pptx: Path
    ) -> None:
        """
        Test split mode - each note saved to separate file.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)

            result = self.run_script(
                script_path, ["-i", str(sample_pptx), "-o", str(output_dir), "--split"]
            )
            assert result.returncode == 0

            # Should create 3 files in a subfolder named after the source file
            source_name = sample_pptx.stem
            subfolder = output_dir / source_name
            output_files = list(subfolder.glob("*.txt"))
            assert len(output_files) == 3

            # Check content of each file
            contents = [f.read_text(encoding="utf-8") for f in sorted(output_files)]
            assert "This is note 1 for slide 1" in contents[0]
            assert "This is note 2 for slide 2" in contents[1]
            assert "With multiple lines" in contents[1]
            assert "kanji" in contents[2]
            assert "memo" not in contents[2]

    def test_extract_notes_split_with_template(
        self, script_path: Path, sample_pptx: Path
    ) -> None:
        """
        Test split mode with template naming pattern.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            template = str(output_dir / "note_%02d.txt")

            result = self.run_script(
                script_path, ["-i", str(sample_pptx), "-o", template, "--split"]
            )
            assert result.returncode == 0

            output_files = list(output_dir.glob("note_*.txt"))
            assert len(output_files) == 3

    def test_extract_notes_empty_presentation(
        self, script_path: Path, empty_pptx: Path
    ) -> None:
        """
        Test extraction from presentation with no notes.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "output.txt"

            result = self.run_script(
                script_path, ["-i", str(empty_pptx), "-o", str(output_file)]
            )
            assert result.returncode == 0
            assert "No notes found" in result.stdout

    def test_extract_notes_glob_pattern(
        self, script_path: Path, sample_pptx: Path
    ) -> None:
        """
        Test extraction using glob pattern for multiple files.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            # Copy sample file to temp dir with different names
            file1 = Path(tmpdir) / "presentation1.pptx"
            file2 = Path(tmpdir) / "presentation2.pptx"
            shutil.copy(sample_pptx, file1)
            shutil.copy(sample_pptx, file2)

            output_dir = Path(tmpdir) / "output"
            output_dir.mkdir()

            pattern = str(Path(tmpdir) / "*.pptx")
            result = self.run_script(
                script_path, ["-i", pattern, "-o", str(output_dir)]
            )
            assert result.returncode == 0

            output_files = list(output_dir.glob("*.txt"))
            assert len(output_files) == 2

    def test_extract_notes_invalid_file(self, script_path: Path) -> None:
        """
        Test handling of non-PowerPoint files.
        """
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
            temp_file = Path(f.name)
            f.write(b"This is not a PowerPoint file")

        try:
            with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
                output_file = Path(f.name)

            result = self.run_script(
                script_path, ["-i", str(temp_file), "-o", str(output_file)]
            )
            assert result.returncode == 0
            assert "Skipping unsupported file type" in result.stdout
        finally:
            if temp_file.exists():
                temp_file.unlink()
            if output_file.exists():
                output_file.unlink()

    def test_extract_notes_nonexistent_file(self, script_path: Path) -> None:
        """
        Test handling of non-existent input file.
        """
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
            output_file = Path(f.name)

        try:
            result = self.run_script(
                script_path, ["-i", "nonexistent.pptx", "-o", str(output_file)]
            )
            assert result.returncode == 0
            assert "No files matching pattern" in result.stdout
        finally:
            if output_file.exists():
                output_file.unlink()

    def test_clean_function(self) -> None:
        """
        Test the clean function directly.
        """
        # Test ruby format cleaning
        assert clean("{kanji|kana}") == "kanji"
        assert clean("{漢字|かな}") == "漢字"
        assert clean("Text {ruby|reading} more") == "Text ruby more"

        # Test memo format cleaning
        assert clean("[memo]") == ""
        assert clean("Text [memo] more") == "Text  more"
        assert clean("[メモ]") == ""

        # Test combined
        assert clean("{kanji|kana} [memo]") == "kanji "

    def test_is_sequence_template(self) -> None:
        """
        Test the is_sequence_template function.
        """
        assert is_sequence_template("output_%03d.txt") is True
        assert is_sequence_template("output_%d.txt") is True
        assert is_sequence_template("output.txt") is False
        assert is_sequence_template("output") is False

    def test_build_split_paths(self) -> None:
        """
        Test the build_split_paths function.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            # Test with directory - creates subfolder named after source file
            # width = len(str(count)) = len("3") = 1, so format is 1, 2, 3 (not 01, 02, 03)
            paths = build_split_paths(tmpdir, "source.pptx", 3, ".txt")
            assert len(paths) == 3
            assert all(p.startswith(tmpdir) for p in paths)
            # The function creates a subfolder "source" and files like "source/1.txt"
            assert paths[0].endswith(os.path.join("source", "1.txt"))
            assert paths[1].endswith(os.path.join("source", "2.txt"))
            assert paths[2].endswith(os.path.join("source", "3.txt"))

            # Test with template
            template = os.path.join(tmpdir, "note_%02d.txt")
            paths = build_split_paths(template, "source.pptx", 3, ".txt")
            assert len(paths) == 3
            assert paths[0].endswith("note_01.txt")
            assert paths[1].endswith("note_02.txt")
            assert paths[2].endswith("note_03.txt")


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
